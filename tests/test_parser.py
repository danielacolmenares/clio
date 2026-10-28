from types import SimpleNamespace

from docx import Document

from clio import config, parser
from tests.pdf_minimo import crea_pdf


def test_pdf_reale_minimo(tmp_path):
    pdf = tmp_path / "atto.pdf"
    pdf.write_bytes(crea_pdf([
        ["Intestazione fissa", "Articolo 1", "Oggetto e finalità"],
        ["Intestazione fissa", "1. Testo del primo paragrafo."],
        ["Intestazione fissa", "2. Testo del secondo paragrafo."],
    ]))
    r = parser.estrai_pdf(pdf)
    assert r.stato == "OK_CON_AVVISI"          # avviso: righe ripetute rimosse
    assert r.pagine == 3
    assert r.testo.startswith("[PAGINA 1]\n")
    assert "[PAGINA 2]" in r.testo and "[PAGINA 3]" in r.testo
    assert "Oggetto e finalità" in r.testo      # carattere accentato preservato
    assert "Intestazione fissa" not in r.testo  # boilerplate rimosso
    assert any("righe ripetute" in a for a in r.avvisi)


def test_pagina_quasi_vuota_genera_avviso(tmp_path):
    pdf = tmp_path / "x.pdf"
    pdf.write_bytes(crea_pdf([["Contenuto sufficiente per la prima pagina."], ["ok"]]))
    r = parser.estrai_pdf(pdf)
    assert any("Pagina 2" in a and "scansionata" in a for a in r.avvisi)


def test_pdf_senza_testo_e_errore(tmp_path, monkeypatch):
    pagina = SimpleNamespace(extract_text=lambda: "", extract_tables=lambda: [])

    class FintoPdf:
        pages = [pagina, pagina]

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda p: FintoPdf())
    r = parser.estrai_pdf(tmp_path / "scansione.pdf")
    assert r.stato == "ERRORE"
    assert "OCR" in r.errore


def test_tabelle_aggiunte_solo_se_richiesto(tmp_path, monkeypatch):
    tabella = [["Fase", "Responsabile"], ["Analisi", "Titolare"]]
    pagina = SimpleNamespace(extract_text=lambda: "Testo della pagina con contenuto.",
                             extract_tables=lambda: [tabella])

    class FintoPdf:
        pages = [pagina]

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    import pdfplumber
    monkeypatch.setattr(pdfplumber, "open", lambda p: FintoPdf())
    con = parser.estrai_pdf(tmp_path / "t.pdf", estrai_tabelle=True)
    senza = parser.estrai_pdf(tmp_path / "t.pdf", estrai_tabelle=False)
    assert "[TABELLA]\nFase: Analisi | Responsabile: Titolare\n[/TABELLA]" in con.testo
    assert "[TABELLA]" not in senza.testo


def test_formatta_tabella():
    assert parser.formatta_tabella([["A", ""], ["1", "2"], [None, None]]) == "A: 1 | Valore: 2"
    assert parser.formatta_tabella([["solo intestazione"]]) == ""


def test_boilerplate_non_rimuove_tag_e_poche_pagine():
    pagine = ["[PAGINA 1]\nX\n[TABELLA]", "[PAGINA 2]\nX\n[TABELLA]", "[PAGINA 3]\nX\n[TABELLA]"]
    pulite, rimosse = parser.rimuovi_boilerplate(pagine)
    assert rimosse == ["X"]
    assert all("[TABELLA]" in p and p.startswith("[PAGINA") for p in pulite)
    assert parser.rimuovi_boilerplate(["a\nX", "b\nX"]) == (["a\nX", "b\nX"], [])


def test_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("1 Scopo")
    doc.add_paragraph("Testo dello scopo.")
    t = doc.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cella A"
    t.cell(0, 1).text = "Cella B"
    percorso = tmp_path / "proc.docx"
    doc.save(percorso)
    r = parser.estrai_docx(percorso)
    assert r.stato == "OK" and r.tipo == "DOCX"
    assert "Testo dello scopo." in r.testo and "Cella B" in r.testo


def test_esegui_scrive_testi_e_report(base):
    kb = config.get_kb("demo")
    (kb.corpus / "atto.pdf").write_bytes(crea_pdf([["Pagina uno con testo sufficiente."]]))
    (kb.corpus / "note.txt").write_text("ignorato")
    kb.testo_estratto.mkdir(parents=True)
    (kb.testo_estratto / "obsoleto.txt").write_text("vecchio")
    assert parser.esegui(kb) == 0
    assert sorted(p.name for p in kb.testo_estratto.iterdir()) == ["atto.txt"]
    report = list((base / "report").glob("parser_report_demo_*.txt"))
    log = list((base / "log").glob("parser_demo_*.log"))
    assert len(report) == 1 and len(log) == 1
    assert "Documenti trovati:      1" in report[0].read_text(encoding="utf-8")
