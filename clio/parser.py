"""Passo 1 di 7 – Parser documenti (PDF/DOCX -> testo).

Cosa fa:
    1. scansiona kb/<nome>/corpus/
    2. estrae il testo da ogni PDF/DOCX
    3. salva ogni documento come .txt in kb/<nome>/testo_estratto/, con un
       marcatore ``[PAGINA N]`` all'inizio di ogni pagina (PDF)
    4. scrive report e log dell'esecuzione

Cosa NON fa: chunking, embedding, Qdrant. Un solo compito, verificabile.

Le righe identiche ripetute su molte pagine (intestazioni e pie' di pagina
fissi) vengono rimosse con un criterio di frequenza, valido per qualunque
documento. Le intestazioni che cambiano a ogni pagina (es. numero di pagina
della Gazzetta ufficiale) non sono catturate da questo criterio: le
rimuove il profilo di chunking specifico (passo 2).
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from .config import KB
from .esecuzione import Esecuzione

ESTENSIONI = {".pdf", ".docx"}

# Sotto questa soglia di caratteri la pagina e' sospetta (scansione senza
# layer di testo, oppure pagina vuota).
MIN_CARATTERI_PAGINA = 20

# Criterio boilerplate: riga corta presente identica in almeno questa
# frazione delle pagine.
SOGLIA_FREQUENZA_BOILERPLATE = 0.5
LUNGHEZZA_MAX_RIGA_BOILERPLATE = 100
# Tag strutturali aggiunti dal parser: mai trattati come rumore.
RIGHE_ESCLUSE_DA_BOILERPLATE = {"[TABELLA]", "[/TABELLA]"}


@dataclass
class RisultatoEstrazione:
    nome_file: str
    tipo: str | None = None          # PDF / DOCX
    pagine: int | None = None        # None per DOCX
    caratteri: int = 0
    stato: str = "OK"                # OK / OK_CON_AVVISI / ERRORE
    avvisi: list[str] = field(default_factory=list)
    errore: str | None = None
    testo: str = ""


def rimuovi_boilerplate(testo_pagine: list[str]) -> tuple[list[str], list[str]]:
    """Rimuove le righe ripetute identiche su molte pagine.
    Restituisce (pagine_pulite, righe_rimosse)."""
    n_pagine = len(testo_pagine)
    if n_pagine < 3:
        # Con poche pagine il criterio di ripetizione non e' affidabile.
        return testo_pagine, []

    righe_per_pagina = [p.split("\n") for p in testo_pagine]
    conteggio: dict[str, int] = {}
    for righe in righe_per_pagina:
        for riga in {r.strip() for r in righe if r.strip()}:
            if riga in RIGHE_ESCLUSE_DA_BOILERPLATE or riga.startswith("[PAGINA "):
                continue
            if len(riga) <= LUNGHEZZA_MAX_RIGA_BOILERPLATE:
                conteggio[riga] = conteggio.get(riga, 0) + 1

    soglia = max(2, int(n_pagine * SOGLIA_FREQUENZA_BOILERPLATE))
    boilerplate = {riga for riga, c in conteggio.items() if c >= soglia}
    pulite = ["\n".join(r for r in righe if r.strip() not in boilerplate) for righe in righe_per_pagina]
    return pulite, sorted(boilerplate)


def formatta_tabella(tabella: list) -> str:
    """Tabella pdfplumber -> righe 'intestazione: valore | ...', usando la
    prima riga come intestazione. Evita colonne appiattite in un flusso
    ambiguo."""
    if not tabella or len(tabella) < 2:
        return ""
    intestazioni = [(c or "").strip() for c in tabella[0]]
    righe_testo = []
    for riga in tabella[1:]:
        celle = [(c or "").strip() for c in riga]
        parti = [
            f"{intest or 'Valore'}: {valore}"
            for intest, valore in zip(intestazioni, celle) if valore
        ]
        if parti:
            righe_testo.append(" | ".join(parti))
    return "\n".join(righe_testo)


def estrai_pdf(path: Path, estrai_tabelle: bool = True, log=None) -> RisultatoEstrazione:
    import pdfplumber  # import locale: il resto del pacchetto non ne dipende

    r = RisultatoEstrazione(path.name, tipo="PDF")
    try:
        with pdfplumber.open(path) as pdf:
            r.pagine = len(pdf.pages)
            testo_pagine = []
            n_tabelle = 0
            for i, pagina in enumerate(pdf.pages, start=1):
                testo = pagina.extract_text() or ""
                if len(testo.strip()) < MIN_CARATTERI_PAGINA:
                    r.avvisi.append(
                        f"Pagina {i}: solo {len(testo.strip())} caratteri estratti. "
                        f"Possibile pagina scansionata o vuota."
                    )
                if estrai_tabelle:
                    try:
                        tabelle = pagina.extract_tables()
                    except Exception as exc:  # pdfplumber puo' fallire su layout anomali
                        tabelle = []
                        r.avvisi.append(f"Pagina {i}: errore estrazione tabelle: {exc}")
                    for tabella in tabelle:
                        testo_tab = formatta_tabella(tabella)
                        if testo_tab:
                            n_tabelle += 1
                            testo += f"\n\n[TABELLA]\n{testo_tab}\n[/TABELLA]"
                testo_pagine.append(f"[PAGINA {i}]\n{testo}")

            if n_tabelle and log:
                log.info("%s - %s tabelle strutturate estratte", path.name, n_tabelle)

            testo_pagine, rimosse = rimuovi_boilerplate(testo_pagine)
            if rimosse:
                r.avvisi.append(
                    f"Rimosse {len(rimosse)} righe ripetute (intestazioni/pie' di pagina): "
                    + " | ".join(rimosse[:5]) + (" ..." if len(rimosse) > 5 else "")
                )
            r.testo = "\n\n".join(testo_pagine)
            r.caratteri = len(r.testo)

        contenuto = r.testo.replace("[PAGINA", "").strip()
        if not any(ch.isalpha() for ch in contenuto):
            r.stato = "ERRORE"
            r.errore = "Nessun testo estratto. Probabile PDF scansionato (serve OCR)."
        elif r.avvisi:
            r.stato = "OK_CON_AVVISI"
    except Exception as exc:
        r.stato = "ERRORE"
        r.errore = str(exc)
    return r


def estrai_docx(path: Path) -> RisultatoEstrazione:
    from docx import Document

    r = RisultatoEstrazione(path.name, tipo="DOCX")
    try:
        doc = Document(path)
        parti = [p.text for p in doc.paragraphs]
        # Anche il testo nelle tabelle, altrimenti andrebbe perso.
        for tabella in doc.tables:
            for riga in tabella.rows:
                for cella in riga.cells:
                    if cella.text.strip():
                        parti.append(cella.text)
        r.testo = "\n".join(parti)
        r.caratteri = len(r.testo)
        if not r.testo.strip():
            r.stato = "ERRORE"
            r.errore = "Nessun testo estratto dal documento DOCX."
    except Exception as exc:
        r.stato = "ERRORE"
        r.errore = str(exc)
    return r


def estrai(path: Path, estrai_tabelle: bool = True, log=None) -> RisultatoEstrazione:
    if path.suffix.lower() == ".pdf":
        return estrai_pdf(path, estrai_tabelle=estrai_tabelle, log=log)
    return estrai_docx(path)


def _report(es: Esecuzione, kb: KB, risultati: list[RisultatoEstrazione]) -> list[str]:
    righe = es.intestazione("PARSER DOCUMENTI - REPORT")
    righe += [
        "RIEPILOGO", "-" * 70,
        f"Documenti trovati:      {len(risultati)}",
        f"PDF:                    {sum(r.tipo == 'PDF' for r in risultati)}",
        f"DOCX:                   {sum(r.tipo == 'DOCX' for r in risultati)}",
        f"Completati con avvisi:  {sum(r.stato == 'OK_CON_AVVISI' for r in risultati)}",
        f"Errori:                 {sum(r.stato == 'ERRORE' for r in risultati)}",
        f"Caratteri estratti:     {sum(r.caratteri for r in risultati)}",
        f"Output:                 {kb.testo_estratto}",
        "", "=" * 70, "DETTAGLIO PER DOCUMENTO", "=" * 70,
    ]
    for r in risultati:
        righe += ["", f"Documento: {r.nome_file}"]
        if r.pagine is not None:
            righe.append(f"Pagine: {r.pagine}")
        righe += [f"Caratteri: {r.caratteri}", f"Stato: {r.stato}"]
        righe += [f"  - {a}" for a in r.avvisi]
        if r.errore:
            righe.append(f"Errore: {r.errore}")
        righe.append("-" * 70)
    return righe


def esegui(kb: KB) -> int:
    """Esegue il passo 1. Restituisce il codice di uscita (0 = ok)."""
    with Esecuzione("parser", kb.nome) as es:
        log = es.logger
        if not kb.corpus.is_dir():
            log.error("Cartella corpus non trovata: %s", kb.corpus)
            return 1
        file_trovati = sorted(
            p for p in kb.corpus.iterdir() if p.is_file() and p.suffix.lower() in ESTENSIONI
        )
        if not file_trovati:
            log.warning("Nessun PDF o DOCX in %s", kb.corpus)
            return 0

        kb.testo_estratto.mkdir(parents=True, exist_ok=True)
        # Rimuove i .txt di documenti non piu' presenti nel corpus, che
        # altrimenti finirebbero nei passi successivi.
        attesi = {p.stem + ".txt" for p in file_trovati}
        for vecchio in kb.testo_estratto.glob("*.txt"):
            if vecchio.name not in attesi:
                vecchio.unlink()
                log.info("Rimosso testo obsoleto: %s", vecchio.name)

        risultati = []
        for path in file_trovati:
            log.info("Elaborazione: %s", path.name)
            r = estrai(path, estrai_tabelle=kb.opzioni["estrai_tabelle"], log=log)
            for a in r.avvisi:
                log.warning("%s - %s", path.name, a)
            if r.stato == "ERRORE":
                log.error("%s - %s", path.name, r.errore)
            else:
                out = kb.testo_estratto / f"{path.stem}.txt"
                out.write_text(r.testo, encoding="utf-8")
                log.info("%s -> %s", path.name, out)
            risultati.append(r)

        report = es.percorso_report("parser_report")
        righe = _report(es, kb, risultati)
        report.write_text("\n".join(righe), encoding="utf-8")
        print("\n".join(righe[:13]))
        print(f"Report: {report}")
        return 2 if any(r.stato == "ERRORE" for r in risultati) else 0
